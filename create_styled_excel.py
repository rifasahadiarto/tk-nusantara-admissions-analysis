import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

csv_file = r"C:\Users\rifas\.gemini\antigravity\scratch\portfolio_tk\tk_nusantara_clean.csv"
excel_file = r"C:\Users\rifas\.gemini\antigravity\scratch\portfolio_tk\tk_nusantara_rapi.xlsx"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Data Siswa Bersih"

# Read CSV
with open(csv_file, mode='r', encoding='utf-8') as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 1):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            # Try convert numeric
            if row_idx > 1:
                # Column 5 (Usia), 7 (Tinggi), 8 (Berat), 9 (Anak Ke)
                if col_idx in [5, 8]:
                    try:
                        cell.value = float(value)
                        cell.number_format = '0.0'
                    except:
                        cell.value = value
                elif col_idx in [7, 9]:
                    try:
                        cell.value = int(value)
                        cell.number_format = '#,##0'
                    except:
                        cell.value = value
                else:
                    cell.value = value
            else:
                # Header friendly names
                friendly_headers = {
                    'ID_Siswa': 'ID Siswa',
                    'Nama_Siswa': 'Nama Lengkap Murid',
                    'Jenis_Kelamin': 'Gender',
                    'Tanggal_Lahir': 'Tgl Lahir',
                    'Usia_Tahun': 'Usia (Thn)',
                    'Pilihan_Kelas': 'Pilihan Kelas',
                    'Tinggi_Badan_cm': 'Tinggi (cm)',
                    'Berat_Badan_kg': 'Berat (kg)',
                    'Anak_Ke': 'Anak Ke-',
                    'Kondisi_Kesehatan_Alergi': 'Catatan Medis & Alergi',
                    'Pekerjaan_Ayah': 'Pekerjaan Ayah',
                    'Pendidikan_Ayah': 'Pendidikan Ayah',
                    'Pekerjaan_Ibu': 'Pekerjaan Ibu',
                    'Pendidikan_Ibu': 'Pendidikan Ibu',
                    'Saluran_Pemasaran': 'Sumber Info Pendaftaran',
                    'Kelurahan_Domisili': 'Desa / Kelurahan',
                    'Kecamatan_Domisili': 'Kecamatan',
                    'Tanggal_Pendaftaran': 'Tgl Daftar'
                }
                cell.value = friendly_headers.get(value, value)

# Styling Definitions
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Deep Navy
zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Format Header
ws.row_dimensions[1].height = 28
for col in range(1, ws.max_column + 1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = align_header
    cell.border = thin_border

# Format Data Rows
for row in range(2, ws.max_row + 1):
    ws.row_dimensions[row].height = 20
    is_even = (row % 2 == 0)
    current_fill = zebra_fill if is_even else white_fill
    
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name="Calibri", size=10)
        cell.fill = current_fill
        cell.border = thin_border
        
        # Alignment rules
        if col in [1, 3, 4, 6, 9, 18]: # ID, Gender, Tgl Lahir, Kelas, Anak Ke, Tgl Daftar
            cell.alignment = align_center
        elif col in [5, 7, 8]: # Usia, Tinggi, Berat
            cell.alignment = align_right
        else:
            cell.alignment = align_left
            
        # Highlight medical alerts
        if col == 10 and cell.value and "Sehat" not in str(cell.value):
            cell.font = Font(name="Calibri", size=10, bold=True, color="9C0006")
            cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

# Auto-adjust column widths with generous padding
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.value:
            max_len = max(max_len, len(str(cell.value)))
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# Specific custom widths
ws.column_dimensions['A'].width = 12 # ID Siswa
ws.column_dimensions['B'].width = 30 # Nama
ws.column_dimensions['C'].width = 14 # Gender
ws.column_dimensions['D'].width = 14 # Tgl Lahir
ws.column_dimensions['E'].width = 12 # Usia
ws.column_dimensions['F'].width = 14 # Kelas
ws.column_dimensions['G'].width = 13 # Tinggi
ws.column_dimensions['H'].width = 13 # Berat
ws.column_dimensions['I'].width = 11 # Anak Ke
ws.column_dimensions['J'].width = 35 # Medis & Alergi
ws.column_dimensions['K'].width = 28 # Pekerjaan Ayah
ws.column_dimensions['L'].width = 22 # Pendidikan Ayah
ws.column_dimensions['M'].width = 24 # Pekerjaan Ibu
ws.column_dimensions['N'].width = 22 # Pendidikan Ibu
ws.column_dimensions['O'].width = 28 # Sumber Info
ws.column_dimensions['P'].width = 18 # Desa
ws.column_dimensions['Q'].width = 18 # Kecamatan
ws.column_dimensions['R'].width = 14 # Tgl Daftar

# Freeze top row
ws.freeze_panes = "A2"

# Enable auto-filter
ws.auto_filter.ref = ws.dimensions

wb.save(excel_file)
print(f"Excel rapi berhasil dibuat: {excel_file}")
